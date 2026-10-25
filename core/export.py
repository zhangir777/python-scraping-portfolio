"""Export records to client-ready Excel / CSV / JSON."""

import csv
import json
from collections.abc import Iterable, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
GBP = '"£"#,##0.00'


def style_sheet(
    ws: Worksheet,
    money_cols: Iterable[str] = (),
    money_format: str = GBP,
    url_cols: Iterable[str] = (),
    percent_cols: Iterable[str] = (),
) -> None:
    """Bold filled header, frozen first row, autofilter, column widths, number formats, links."""
    headers = [c.value for c in ws[1]]
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    money, urls, pcts = set(money_cols), set(url_cols), set(percent_cols)
    for idx, name in enumerate(headers, start=1):
        width = len(str(name))
        for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            if name in money:
                cell.number_format = money_format
            elif name in pcts:
                cell.number_format = "0.0%"
            elif name in urls and cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = LINK_FONT
            width = max(width, len(str(cell.value or "")))
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 60)


def _cell(value):
    if isinstance(value, list):
        return ", ".join(map(str, value))
    if value is None or isinstance(value, (int, float, str, bool)):
        return value
    return str(value)


def write_rows(ws: Worksheet, rows: Sequence[dict], columns: Sequence[str]) -> None:
    ws.append(list(columns))
    for row in rows:
        ws.append([_cell(row.get(c)) for c in columns])


def add_summary_sheet(ws: Worksheet, summary: dict) -> None:
    """summary: {label: value}; a dict value becomes its own small table."""
    ws.append(["Metric", "Value"])
    for key, value in summary.items():
        if isinstance(value, dict):
            ws.append([])
            ws.append([key])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for k, v in value.items():
                ws.append([k, v])
        else:
            ws.append([key, value])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24


def export_xlsx(
    rows: Sequence[dict],
    path: Path,
    columns: Sequence[str],
    summary: dict | None = None,
    sheet_title: str = "Data",
    **style,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    write_rows(ws, rows, columns)
    style_sheet(ws, **style)
    if summary:
        add_summary_sheet(wb.create_sheet("Summary"), summary)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_csv(rows: Sequence[dict], path: Path, columns: Sequence[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so Excel on Windows detects the encoding
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({c: _cell(row.get(c)) for c in columns})
    return path


def export_json(rows: Sequence[dict], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(list(rows), ensure_ascii=False, indent=2, default=str), "utf-8")
    return path


def export_all(
    rows: Sequence[dict], base: Path, formats: Iterable[str], columns: Sequence[str], **xlsx_kwargs
) -> list[Path]:
    out = []
    for fmt in formats:
        fmt = fmt.strip().lower()
        target = base.with_suffix(f".{fmt}")
        if fmt == "xlsx":
            out.append(export_xlsx(rows, target, columns, **xlsx_kwargs))
        elif fmt == "csv":
            out.append(export_csv(rows, target, columns))
        elif fmt == "json":
            out.append(export_json(rows, target))
        else:
            raise ValueError(f"Unknown format: {fmt}")
    return out

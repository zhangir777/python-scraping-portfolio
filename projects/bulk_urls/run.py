"""Prices for a long list of product URLs on many sites -> Excel with Prices, Failed and Summary.

uv run bulk-prices samples/bulk_urls.csv --out samples/bulk_prices
Sites run in parallel, each site has its own polite queue (1 request/s by default).
Every result is written to <out>.csv at once, so a crash or Ctrl+C loses nothing:
run the same command again and it continues with the URLs not done yet.
"""

import csv
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

import httpx
import typer
import yaml
from openpyxl import Workbook, load_workbook

from core.export import add_summary_sheet, style_sheet, write_rows
from core.http import PoliteClient, RobotsDisallowed
from core.logging import setup_logging

from .extract import extract_price

COLUMNS = ["url", "site", "title", "price", "currency", "method", "status", "error", "checked_at"]
SITES = Path(__file__).with_name("sites.yaml")
app = typer.Typer(add_completion=False)


def read_urls(path: Path) -> list[str]:
    """URLs from the `url` column of a CSV/XLSX (or the first column), duplicates removed."""
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        rows = list(load_workbook(path, read_only=True).active.iter_rows(values_only=True))
        header = [str(c or "").strip().lower() for c in rows[0]]
        col = header.index("url") if "url" in header else 0
        values = [r[col] for r in rows[1:]]
    else:
        with path.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            col = "url" if "url" in (reader.fieldnames or []) else reader.fieldnames[0]
            values = [r[col] for r in reader]
    urls = [str(v).strip() for v in values if v and str(v).strip().startswith("http")]
    return list(dict.fromkeys(urls))


def read_done(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def check_url(client: PoliteClient, url: str, selectors: dict | None) -> dict:
    row = {"url": url, "site": urlsplit(url).netloc, "checked_at": datetime.now().isoformat("T")}
    try:
        found = extract_price(client.get_text(url), selectors)
    except RobotsDisallowed:
        return {**row, "status": "skipped", "error": "blocked by robots.txt"}
    except httpx.HTTPStatusError as exc:
        return {**row, "status": "failed", "error": f"HTTP {exc.response.status_code}"}
    except httpx.HTTPError as exc:
        return {**row, "status": "failed", "error": str(exc)[:200]}
    if found["price"] is None:
        return {**row, **found, "status": "failed", "error": "price not found"}
    return {**row, **found, "status": "ok", "error": ""}


def run(urls: list[str], out_csv: Path, sites: dict, workers: int, delay: float, log) -> None:
    by_site: dict[str, list[str]] = defaultdict(list)
    for url in urls:
        by_site[urlsplit(url).netloc].append(url)
    lock = threading.Lock()
    new_file = not out_csv.exists()
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    with out_csv.open("a", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, COLUMNS, extrasaction="ignore")
        if new_file:
            writer.writeheader()

        def work(site: str) -> None:
            client = PoliteClient(delay=delay)
            try:
                for url in by_site[site]:
                    row = check_url(client, url, sites.get(site))
                    with lock:
                        writer.writerow(row)
                        f.flush()
                    log.info("%s %s %s", row["status"], row.get("price") or "", url)
            finally:
                client.close()

        with ThreadPoolExecutor(max_workers=workers) as pool:
            list(pool.map(work, by_site))


def build_report(rows: list[dict], path: Path) -> Path:
    ok = [{**r, "price": float(r["price"])} for r in rows if r["status"] == "ok"]
    bad = [r for r in rows if r["status"] != "ok"]
    per_site = Counter(r["site"] for r in rows)
    ok_site = Counter(r["site"] for r in ok)
    times = sorted(r["checked_at"][:16].replace("T", " ") for r in rows)
    summary = {
        "URLs": len(rows),
        "Prices found": len(ok),
        "Failed / skipped": len(bad),
        "Success rate": f"{len(ok) / len(rows):.0%}" if rows else "0%",
        "Checked": f"{times[0]} to {times[-1]}" if times else "",
        "Created": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Per site (ok / total)": {s: f"{ok_site[s]} / {n}" for s, n in per_site.items()},
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    write_rows(ws, ok, ["site", "title", "price", "currency", "method", "url", "checked_at"])
    style_sheet(ws, money_cols=["price"], url_cols=["url"])
    failed = wb.create_sheet("Failed")
    write_rows(failed, bad, ["site", "url", "status", "error", "checked_at"])
    style_sheet(failed, url_cols=["url"])
    add_summary_sheet(wb.create_sheet("Summary", 0), summary)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@app.command()
def main(
    input: Path = typer.Argument(..., help="CSV or Excel with a `url` column"),
    out: Path = typer.Option(Path("out/bulk_prices"), help="Output path without extension"),
    workers: int = typer.Option(8, help="Sites processed at the same time"),
    delay: float = typer.Option(1.0, help="Seconds between requests to the same site"),
    retry_failed: bool = typer.Option(False, help="Check failed URLs again"),
) -> None:
    log = setup_logging("bulk")
    started = time.monotonic()
    sites = yaml.safe_load(SITES.read_text(encoding="utf-8")) or {}
    out_csv = out.with_suffix(".csv")
    done = read_done(out_csv)
    if retry_failed:
        done = [r for r in done if r["status"] == "ok"]
        with out_csv.open("w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, COLUMNS)
            writer.writeheader()
            writer.writerows(done)
    urls = read_urls(input)
    seen = {r["url"] for r in done}
    todo = [u for u in urls if u not in seen]
    log.info("%d URLs, %d already done, %d to check", len(urls), len(seen), len(todo))
    run(todo, out_csv, sites, workers, delay, log)

    wanted = set(urls)
    rows = [r for r in read_done(out_csv) if r["url"] in wanted]
    report = build_report(rows, out.with_suffix(".xlsx"))
    log.info("This run: %.0fs", time.monotonic() - started)
    log.info("Saved %s and %s", report, out_csv)


if __name__ == "__main__":
    app()

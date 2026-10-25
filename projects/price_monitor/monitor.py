"""Compare two snapshots and write changes.xlsx; optional Telegram summary.

uv run price-monitor tests/fixtures/books_yesterday.csv samples/books.csv --out samples/changes.xlsx
"""

import csv
import os
from pathlib import Path

import httpx
import typer
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from core.export import add_summary_sheet, style_sheet, write_rows
from core.logging import setup_logging

from .compare import Changes, compare

RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="C6EFCE")
app = typer.Typer(add_completion=False)


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_report(ch: Changes, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Price changes"
    cols = ["upc", "title", "old_price", "new_price", "diff", "diff_pct", "product_url"]
    write_rows(ws, ch.price, cols)
    style_sheet(
        ws,
        money_cols=["old_price", "new_price", "diff"],
        percent_cols=["diff_pct"],
        url_cols=["product_url"],
    )
    for row in ws.iter_rows(min_row=2):
        fill = RED if row[4].value > 0 else GREEN  # price up = red, down = green
        for cell in row[2:6]:
            cell.fill = fill

    sheets = [
        (
            "Stock changes",
            ch.stock,
            ["upc", "title", "was_in_stock", "now_in_stock", "product_url"],
        ),
        ("New products", ch.added, ["upc", "title", "price_gbp", "in_stock", "product_url"]),
        ("Removed products", ch.removed, ["upc", "title", "price_gbp", "product_url"]),
    ]
    for title, rows, cols in sheets:
        s = wb.create_sheet(title)
        rows = [{**r, "price_gbp": float(r["price_gbp"])} if "price_gbp" in r else r for r in rows]
        write_rows(s, rows, cols)
        style_sheet(s, money_cols=["price_gbp"], url_cols=["product_url"])
    add_summary_sheet(wb.create_sheet("Summary", 0), ch.summary())
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def notify(text: str, log) -> None:
    token, chat = os.getenv("TELEGRAM_BOT_TOKEN"), os.getenv("TELEGRAM_CHAT_ID")
    if not token or not chat:
        log.warning("TELEGRAM_BOT_TOKEN / TELEGRAM_CHAT_ID not set, skipping notification")
        return
    resp = httpx.post(
        f"https://api.telegram.org/bot{token}/sendMessage",
        json={"chat_id": chat, "text": text},
        timeout=15,
    )
    log.info("Telegram notification: %s", "sent" if resp.is_success else resp.status_code)


@app.command()
def main(
    old: Path = typer.Argument(..., help="Previous snapshot CSV"),
    new: Path = typer.Argument(..., help="Current snapshot CSV"),
    out: Path = typer.Option(Path("samples/changes.xlsx")),
    send: bool = typer.Option(False, "--notify", help="Send summary to Telegram"),
) -> None:
    log = setup_logging("monitor")
    load_dotenv()
    ch = compare(read_csv(old), read_csv(new))
    write_report(ch, out)
    text = "Price monitor\n" + "\n".join(f"{k}: {v}" for k, v in ch.summary().items())
    log.info(text.replace("\n", " | "))
    log.info("Saved %s", out)
    if send:
        notify(text, log)


if __name__ == "__main__":
    app()

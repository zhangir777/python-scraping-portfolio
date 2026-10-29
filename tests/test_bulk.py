from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from core.http import PoliteClient
from projects.bulk_urls.extract import extract_price, parse_price
from projects.bulk_urls.run import build_report, check_url, read_done, read_urls, run

FIX = Path(__file__).parent / "fixtures"


@pytest.mark.parametrize(
    "raw, want",
    [
        ("£51.77", 51.77),
        ("$1,299.50", 1299.5),
        ("1 299,50 €", 1299.5),
        ("1,299", 1299.0),
        ("63.00", 63.0),
        (12, 12.0),
        ("n/a", None),
    ],
)
def test_parse_price(raw, want):
    assert parse_price(raw) == want


def test_json_ld_product():
    found = extract_price((FIX / "product_jsonld.html").read_text(encoding="utf-8"))
    assert found == {"title": "Bulbasaur", "price": 63.0, "currency": "GBP", "method": "json-ld"}


def test_meta_and_selector():
    meta = '<h1>Lamp</h1><meta property="product:price:amount" content="19.90">'
    assert extract_price(meta)["price"] == 19.9
    book = (FIX / "book_detail.html").read_text(encoding="utf-8")
    found = extract_price(book, {"price": "p.price_color", "title": "h1", "currency": "GBP"})
    assert found["method"] == "selector" and found["price"] > 0
    assert extract_price("<h1>Nothing</h1>")["price"] is None


def test_read_urls_dedupes(tmp_path):
    p = tmp_path / "in.csv"
    p.write_text("url\nhttps://a.com/1\nhttps://a.com/1\n\nnot a url\nhttps://b.com/2\n")
    assert read_urls(p) == ["https://a.com/1", "https://b.com/2"]


def fake_client(pages: dict[str, tuple[int, str]]) -> PoliteClient:
    def handler(req: httpx.Request) -> httpx.Response:
        if req.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nDisallow: /private/")
        code, body = pages.get(str(req.url), (404, ""))
        return httpx.Response(code, text=body)

    c = PoliteClient(delay=0, retries=0)
    c.client = httpx.Client(transport=httpx.MockTransport(handler))
    return c


def test_check_url_statuses():
    c = fake_client({"https://a.com/ok": (200, '<meta itemprop="price" content="5">')})
    assert check_url(c, "https://a.com/ok", None)["status"] == "ok"
    assert check_url(c, "https://a.com/missing", None)["error"] == "HTTP 404"
    assert check_url(c, "https://a.com/private/x", None)["status"] == "skipped"


def test_resume_and_report(tmp_path, monkeypatch, caplog):
    pages = {
        f"https://s{i}.com/p": (200, f'<meta itemprop="price" content="{i}">') for i in range(4)
    }
    monkeypatch.setattr("projects.bulk_urls.run.PoliteClient", lambda delay: fake_client(pages))
    out = tmp_path / "res.csv"
    import logging

    log = logging.getLogger("t")
    run(list(pages)[:2], out, {}, 2, 0, log)
    run(list(pages)[2:] + ["https://s9.com/p"], out, {}, 2, 0, log)  # second run appends
    rows = read_done(out)
    assert len(rows) == 5 and sum(r["status"] == "ok" for r in rows) == 4
    wb = load_workbook(build_report(rows, tmp_path / "r.xlsx"))
    assert wb.sheetnames == ["Summary", "Prices", "Failed"]
    assert wb["Failed"].max_row == 2

from openpyxl import load_workbook

from core.export import export_all
from projects.price_monitor.compare import compare
from projects.price_monitor.monitor import write_report

OLD = [
    {"upc": "1", "title": "A", "price_gbp": "10.00", "in_stock": "True"},
    {"upc": "2", "title": "B", "price_gbp": "20.00", "in_stock": "True"},
    {"upc": "3", "title": "C", "price_gbp": "30.00", "in_stock": "False"},
]
NEW = [
    {"upc": "1", "title": "A", "price_gbp": "12.00", "in_stock": "True"},
    {"upc": "3", "title": "C", "price_gbp": "27.00", "in_stock": "True"},
    {"upc": "4", "title": "D", "price_gbp": "5.00", "in_stock": "True"},
]


def test_compare_snapshots():
    ch = compare(OLD, NEW)
    assert [r["upc"] for r in ch.added] == ["4"]
    assert [r["upc"] for r in ch.removed] == ["2"]
    by_upc = {c["upc"]: c for c in ch.price}
    assert by_upc["1"]["diff"] == 2.0 and by_upc["1"]["diff_pct"] == 0.2
    assert by_upc["3"]["diff"] == -3.0 and by_upc["3"]["diff_pct"] == -0.1
    assert [(s["upc"], s["now_in_stock"]) for s in ch.stock] == [("3", True)]
    assert ch.summary()["  price up"] == 1


def test_no_changes():
    ch = compare(OLD, OLD)
    assert not (ch.added or ch.removed or ch.price or ch.stock)


def test_changes_report(tmp_path):
    path = write_report(compare(OLD, NEW), tmp_path / "changes.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "Summary"
    ws = wb["Price changes"]
    assert ws["A1"].value == "upc" and ws.max_row == 3
    fills = {ws.cell(r, 1).value: ws.cell(r, 5).fill.fgColor.rgb for r in (2, 3)}
    assert fills["1"].endswith("F8CBAD")  # up = red
    assert fills["3"].endswith("C6EFCE")  # down = green


def test_export_formats(tmp_path):
    rows = [{"title": "A", "price_gbp": 1.5, "product_url": "https://x.test/a", "tags": ["x", "y"]}]
    cols = ["title", "price_gbp", "product_url", "tags"]
    files = export_all(
        rows,
        tmp_path / "out",
        ["xlsx", "csv", "json"],
        cols,
        summary={"Records": 1},
        money_cols=["price_gbp"],
        url_cols=["product_url"],
    )
    assert all(f.exists() for f in files)
    wb = load_workbook(files[0])
    ws = wb.active
    assert [c.value for c in ws[1]] == cols
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold
    assert ws["C2"].hyperlink.target == "https://x.test/a"
    assert ws["D2"].value == "x, y"
    assert "Summary" in wb.sheetnames
    assert files[1].read_text(encoding="utf-8-sig").splitlines()[0] == ",".join(cols)

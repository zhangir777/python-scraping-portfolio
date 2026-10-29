import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from projects.inventory_tracker.tracker import previous_snapshot, track

FIX = Path(__file__).parent / "fixtures"
SAMPLES = Path(__file__).parents[1] / "samples"


def fake_scrape(src: Path):
    def scrape(out: Path) -> None:
        shutil.copy(src, out.with_suffix(".csv"))
        out.with_suffix(".xlsx").write_bytes(b"")

    return scrape


def test_previous_snapshot(tmp_path):
    for name in ["2026-10-01_0800.csv", "2026-10-02_0800.csv", "2026-10-03_0800.csv"]:
        (tmp_path / name).write_text("")
    cur = tmp_path / "2026-10-03_0800.csv"
    assert previous_snapshot(tmp_path, cur).name == "2026-10-02_0800.csv"
    assert previous_snapshot(tmp_path, tmp_path / "2026-10-01_0800.csv") is None


def test_two_runs_give_change_report(tmp_path):
    report, summary = track(
        tmp_path, fake_scrape(FIX / "books_yesterday.csv"), datetime(2026, 10, 1, 8)
    )
    assert "First run" in summary and report.name == "catalog_latest.xlsx"

    report, summary = track(tmp_path, fake_scrape(SAMPLES / "books.csv"), datetime(2026, 10, 2, 8))
    assert summary["Compared with"] == "2026-10-01_0800"
    assert summary["Price changes"] > 0 and summary["New products"] > 0
    assert report.exists()
    assert "Price changes" in load_workbook(report).sheetnames
